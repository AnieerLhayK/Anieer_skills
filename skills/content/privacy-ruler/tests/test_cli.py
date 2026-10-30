"""Synthetic tests at the public CLI seam; run with --basetemp in staging."""
import json
import subprocess
import sys
from pathlib import Path
import pytest

CLI = Path(__file__).parents[1] / "scripts" / "privacy_ruler.py"


def call(*args):
    result = subprocess.run([sys.executable, str(CLI), *map(str, args)], capture_output=True, text=True)
    return result, json.loads(result.stdout)


def prepare(tmp_path, inputs, rules=None, **settings):
    request = tmp_path / "request.json"
    request.write_text(json.dumps({"inputs": [{"path": str(p)} for p in inputs], "rules": rules or [], **settings}), encoding="utf-8")
    private = tmp_path / "private"
    result, _ = call("inspect", "--request", request, "--private-dir", private)
    assert result.returncode == 0
    path = private / "plan.json"
    plan = json.loads(path.read_text(encoding="utf-8"))
    for item in plan["files"]:
        item["semantic_review"] = {"status": "complete", "reviewer": "codex", "coverage": "all_units", "risks": [], "losses": []}
    path.write_text(json.dumps(plan), encoding="utf-8")
    return private, path


@pytest.mark.parametrize("suffix,text", [(".md", "Normal text CanaryValue"), (".json", '{"CanaryValue":{"id":123456,"text":"Normal text"}}'), (".jsonl", '{"text":"CanaryValue","id":123456}\n'), (".csv", 'name,id\nCanaryValue,123456\n'), (".tsv", 'name\tid\nCanaryValue\t123456\n')])
def test_supported_structured_formats_and_scalar_identity(tmp_path, suffix, text):
    source = tmp_path / ("source" + suffix)
    source.write_text(text, encoding="utf-8")
    private, plan = prepare(tmp_path, [source], [{"values": ["CanaryValue", "123456"], "category": "person"}])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 0, report
    output = (tmp_path / "share" / ("file-0001" + suffix)).read_text(encoding="utf-8")
    assert "CanaryValue" not in output and "123456" not in output
    assert "[PERSON_001]" in output
    assert source.read_text(encoding="utf-8") == text


def test_source_change_and_duplicate_keys_fail_closed_without_raw_diagnostic(tmp_path):
    a, b = tmp_path / "a.txt", tmp_path / "b.json"
    a.write_text("CanaryValue", encoding="utf-8")
    b.write_text('{"duplicate":"CanaryValue","duplicate":"other"}', encoding="utf-8")
    private, plan = prepare(tmp_path, [a, b], [{"values": ["CanaryValue"]}])
    a.write_text("changed CanaryValue", encoding="utf-8")
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 2 and not (tmp_path / "share").exists()
    assert "source_changed" in result.stdout
    assert "CanaryValue" not in result.stdout + result.stderr


def test_docx_cross_runs_hidden_history_and_visual_gate(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Mira")
    p.add_run("Canary")
    doc.add_paragraph("Ordinary prose stays.")
    doc.sections[0].header.paragraphs[0].text = "HeaderCanary"
    doc.sections[0].footer.paragraphs[0].text = "FooterCanary"
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "TableCanary"
    hidden = doc.add_paragraph().add_run("HiddenCanary")
    hidden.font.hidden = True
    doc.core_properties.author = "MetadataCanary"
    deleted = OxmlElement("w:del")
    r, t = OxmlElement("w:r"), OxmlElement("w:delText")
    t.text = "HistoryCanary"
    r.append(t)
    deleted.append(r)
    p._p.append(deleted)
    source = tmp_path / "source.docx"
    doc.save(source)
    private, plan = prepare(tmp_path, [source], [{"values": ["MiraCanary", "HiddenCanary", "MetadataCanary", "HistoryCanary", "HeaderCanary", "FooterCanary", "TableCanary"]}])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 2 and not (tmp_path / "share").exists()
    run = json.loads((private / "run.json").read_text(encoding="utf-8"))
    item = run["files"][0]
    assert item["state"] == "awaiting_visual_review", item
    import zipfile
    with zipfile.ZipFile(item["candidate"]) as z:
        payload = b"\n".join(z.read(n) for n in z.namelist())
    for value in [b"MiraCanary", b"HiddenCanary", b"MetadataCanary", b"HistoryCanary", b"HeaderCanary", b"FooterCanary", b"TableCanary"]:
        assert value not in payload
    assert "Ordinary prose stays." in "\n".join(p.text for p in Document(item["candidate"]).paragraphs)
    # Synthetic receipt tests the gate; independent forward testing inspects pages.
    item["visual_review"] = {"status": "complete", "reviewer": "codex", "sha256": item["sha256"], "pages": list(range(1, item["pages"] + 1))}
    (private / "run.json").write_text(json.dumps(run), encoding="utf-8")
    result, report = call("verify", "--run", private / "run.json")
    assert result.returncode == 0
    assert report["files"][0]["state"] == "delivered"
    assert report["files"][0]["code"] == "verified_pending_human_review"


def test_unsupported_embedded_image_blocks_only_its_file_and_reports_partial(tmp_path):
    import pymupdf
    from PIL import Image
    from docx import Document
    image = tmp_path / "picture.png"
    Image.new("RGB", (30, 30), (255, 0, 0)).save(image)
    doc = Document()
    doc.add_paragraph("CanaryValue")
    doc.add_picture(str(image))
    word = tmp_path / "with-image.docx"
    doc.save(word)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((30, 30), "CanaryValue")
    page.insert_image(pymupdf.Rect(50, 50, 100, 100), filename=str(image))
    pdf_path = tmp_path / "with-image.pdf"
    pdf.save(pdf_path)
    pdf.close()
    text = tmp_path / "plain.txt"
    text.write_text("CanaryValue normal content", encoding="utf-8")
    private, plan = prepare(tmp_path, [text, word, pdf_path, image], [{"values": ["CanaryValue"]}])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 2 and report["status"] == "partial"
    assert [i["state"] for i in report["files"]] == ["delivered", "blocked", "blocked", "blocked"]
    assert sorted(p.name for p in (tmp_path / "share").iterdir()) == ["file-0001.txt", "report.json"]
    assert "CanaryValue" not in result.stdout + result.stderr


def test_mapping_reuse_is_explicit_and_aliases_stay_consistent(tmp_path):
    a = tmp_path / "a.txt"
    a.write_text("CanaryPerson AliasCanary", encoding="utf-8")
    private, plan = prepare(tmp_path, [a], [{"values": ["CanaryPerson", "AliasCanary"], "category": "person"}])
    call("apply", "--plan", plan, "--out", tmp_path / "share")
    request = tmp_path / "reuse.json"
    request.write_text(json.dumps({"inputs": [{"path": str(a)}], "mapping_path": str(private / "mapping.json")}), encoding="utf-8")
    result, _ = call("inspect", "--request", request, "--private-dir", tmp_path / "second-private")
    assert result.returncode == 0
    first = json.loads((private / "mapping.json").read_text(encoding="utf-8"))
    second = json.loads((tmp_path / "second-private" / "mapping.json").read_text(encoding="utf-8"))
    assert first == second
    assert len({e["new"] for e in first["replacements"]}) == 1


def test_key_collision_and_invalid_cli_never_leak_raw_values(tmp_path):
    source = tmp_path / "data.json"
    source.write_text('{"CanaryOne":1,"CanaryTwo":2}', encoding="utf-8")
    private, plan = prepare(tmp_path, [source], [{"values": ["CanaryOne", "CanaryTwo"]}])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 2
    assert report["files"][0]["code"] == "duplicate_or_colliding_key"
    result, report = call("inspect", "--CanarySecret")
    assert result.returncode == 2 and "CanarySecret" not in result.stdout + result.stderr


def test_labeled_credentials_protect_repeated_bare_values_without_reusable_secret(tmp_path):
    source = tmp_path / "credentials.txt"
    source.write_text("token: synthetic-secret-789; repeated synthetic-secret-789. Normal prose.", encoding="utf-8")
    private, plan = prepare(tmp_path, [source])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 0
    assert "synthetic-secret-789" not in (tmp_path / "share" / "file-0001.txt").read_text(encoding="utf-8")
    assert "synthetic-secret-789" not in (private / "mapping.json").read_text(encoding="utf-8")


def test_reused_mapping_with_removed_credential_does_not_merge_new_identity(tmp_path):
    source = tmp_path / "first.txt"
    source.write_text("Alice Bob", encoding="utf-8")
    private, plan = prepare(tmp_path, [source], [{"category": "person", "values": ["Alice"]}, {"category": "credential", "values": ["PrivateCredential"]}, {"category": "person", "values": ["Bob"]}])
    call("apply", "--plan", plan, "--out", tmp_path / "first-share")
    request = tmp_path / "reuse.json"
    request.write_text(json.dumps({"inputs": [{"text": "Bob Carol"}], "mapping_path": str(private / "mapping.json"), "rules": [{"category": "person", "values": ["Carol"]}]}), encoding="utf-8")
    second = tmp_path / "second-private"
    call("inspect", "--request", request, "--private-dir", second)
    reused = json.loads((second / "mapping.json").read_text(encoding="utf-8"))
    identities = {e["old"]: e["new"] for e in reused["replacements"]}
    assert identities["Alice"] != identities["Bob"] != identities["Carol"]


def test_pdf_physically_removes_text_and_blocks_images(tmp_path):
    import pymupdf
    pdf = pymupdf.open()
    p = pdf.new_page()
    p.insert_text((50, 50), "MiraCanary Ordinary prose stays.")
    p.insert_text((50, 90), "HiddenCanary", render_mode=3)
    p.add_text_annot((100, 150), "CommentCanary")
    p.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 45, 200, 55), "uri": "https://example.invalid/LinkCanary"})
    pdf.set_toc([[1, "OutlineCanary", 1]])
    pdf.set_metadata({"author": "MetadataCanary"})
    source = tmp_path / "source.pdf"
    pdf.save(source)
    pdf.close()
    private, plan = prepare(tmp_path, [source], [{"values": ["MiraCanary", "MetadataCanary", "HiddenCanary", "CommentCanary", "LinkCanary", "OutlineCanary"], "category": "person"}])
    _, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    run = json.loads((private / "run.json").read_text(encoding="utf-8"))
    item = run["files"][0]
    assert item["state"] == "awaiting_visual_review", item
    with pymupdf.open(item["candidate"]) as doc:
        assert "MiraCanary" not in doc[0].get_text()
        assert "Ordinary prose stays." in doc[0].get_text()
        assert doc[0].get_text().count("[PERSON_001]") == 1, "Hidden source text must not become visible pseudonym text"
        assert not doc.metadata.get("author")
        assert not any(b"MiraCanary" in doc.xref_stream(i) for i in range(1, doc.xref_length()) if doc.xref_is_stream(i))
    assert not (tmp_path / "share").exists()


def test_pdf_page_attachment_blocks_instead_of_silent_removal(tmp_path):
    import pymupdf
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "Ordinary visible text.")
    page.add_file_annot((50, 100), b"Synthetic attached private data", "attachment.txt")
    source = tmp_path / "attachment.pdf"
    pdf.save(source)
    pdf.close()
    private, plan = prepare(tmp_path, [source])
    _, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert report["status"] == "blocked"
    assert report["files"][0]["code"] == "unsupported_document_object"
    assert not (tmp_path / "share").exists()


def test_docx_footnotes_endnotes_and_textbox_are_in_the_same_coverage_boundary(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from lxml import etree
    import zipfile
    import xml.etree.ElementTree as ET
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    doc = Document()
    doc.add_paragraph("Ordinary prose stays.")
    pict = OxmlElement("w:pict")
    rect = etree.SubElement(pict, "{urn:schemas-microsoft-com:vml}rect", nsmap={"v": "urn:schemas-microsoft-com:vml"})
    rect.set("style", "width:180pt;height:40pt")
    box = etree.SubElement(rect, "{urn:schemas-microsoft-com:vml}textbox")
    content = etree.SubElement(box, "{" + w + "}txbxContent")
    p = etree.SubElement(content, "{" + w + "}p")
    r = etree.SubElement(p, "{" + w + "}r")
    etree.SubElement(r, "{" + w + "}t").text = "TextboxCanary"
    doc.add_paragraph().add_run()._r.append(pict)
    for kind in ("footnote", "endnote"):
        ref = OxmlElement("w:" + kind + "Reference")
        ref.set("{" + w + "}id", "1")
        doc.paragraphs[0].add_run()._r.append(ref)
    source = tmp_path / "notes.docx"
    doc.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
    types = ET.fromstring(parts["[Content_Types].xml"])
    for i, (kind, value) in enumerate((("footnote", "FootnoteCanary"), ("endnote", "EndnoteCanary")), 80):
        root = ET.Element("{" + w + "}" + kind + "s")
        note = ET.SubElement(root, "{" + w + "}" + kind, {"{" + w + "}id": "1"})
        p = ET.SubElement(note, "{" + w + "}p")
        r = ET.SubElement(p, "{" + w + "}r")
        ET.SubElement(r, "{" + w + "}t").text = value
        parts["word/" + kind + "s.xml"] = ET.tostring(root)
        ET.SubElement(rels, "{" + rel_ns + "}Relationship", {"Id": f"rId{i}", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/" + kind + "s", "Target": kind + "s.xml"})
        ET.SubElement(types, "{" + ct_ns + "}Override", {"PartName": "/word/" + kind + "s.xml", "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml." + kind + "s+xml"})
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
    parts["[Content_Types].xml"] = ET.tostring(types)
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for n, data in parts.items():
            archive.writestr(n, data)
    private, plan = prepare(tmp_path, [source], [{"values": ["FootnoteCanary", "EndnoteCanary", "TextboxCanary"]}])
    _, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    run = json.loads((private / "run.json").read_text(encoding="utf-8"))
    assert run["files"][0]["state"] == "awaiting_visual_review", run
    with zipfile.ZipFile(run["files"][0]["candidate"]) as archive:
        payload = b"\n".join(archive.read(n) for n in archive.namelist())
    assert all(value not in payload for value in (b"FootnoteCanary", b"EndnoteCanary", b"TextboxCanary"))


def test_public_escalation_semantic_gate_and_consistent_batch(tmp_path):
    sources = []
    for i in range(2):
        p = tmp_path / f"private-name-{i}.txt"
        p.write_text("Contact MiraCanary at canary@example.invalid. Ordinary prose stays.", encoding="utf-8")
        sources.append({"path": str(p)})
    request = tmp_path / "request.json"
    request.write_text(json.dumps({"inputs": sources, "level": "low", "destination": "public",
                                   "rules": [{"category": "person", "values": ["MiraCanary"]}]}), encoding="utf-8")
    private, out = tmp_path / "private", tmp_path / "share"
    result, info = call("inspect", "--request", request, "--private-dir", private)
    assert result.returncode == 0
    assert info["effective_level"] == "high"
    plan_path = private / "plan.json"
    _, info = call("apply", "--plan", plan_path, "--out", out)
    assert info["status"] == "blocked"
    assert not out.exists()
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    for item in plan["files"]:
        item["semantic_review"] = {"status": "complete", "reviewer": "codex", "coverage": "all_units", "risks": [], "losses": []}
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    result, info = call("apply", "--plan", plan_path, "--out", out)
    assert result.returncode == 0, info
    assert info["status"] == "processed_pending_human_review"
    texts = [p.read_text(encoding="utf-8") for p in out.glob("*.txt")]
    assert len(texts) == 2 and texts[0] == texts[1]
    assert "Ordinary prose stays." in texts[0]
    assert "MiraCanary" not in texts[0] and "canary@example.invalid" not in texts[0]
    assert "private-name" not in " ".join(p.name for p in out.iterdir())
    assert "MiraCanary" not in (out / "report.json").read_text(encoding="utf-8")
    assert sources[0]["path"] not in result.stdout + result.stderr


def test_docx_internal_part_name_cannot_carry_protected_identifier(tmp_path):
    from docx import Document
    import zipfile
    source = tmp_path / "package.docx"
    doc = Document()
    doc.add_paragraph("Ordinary prose stays.")
    doc.save(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("word/PartNameCanary.xml", b"<extra>Ordinary metadata.</extra>")
    _, plan = prepare(tmp_path, [source], [{"values": ["PartNameCanary"]}])
    result, report = call("apply", "--plan", plan, "--out", tmp_path / "share")
    assert result.returncode == 2
    assert report["status"] == "blocked"
    assert report["files"][0]["code"] == "protected_value_remaining"
    assert "PartNameCanary" not in result.stdout + result.stderr
    assert not (tmp_path / "share").exists()
