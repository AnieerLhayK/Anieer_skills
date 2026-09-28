from __future__ import annotations

import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

from docx import Document
from PIL import Image

SKILL_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_DIR / "scripts"))

from mardocx_config import load_config
from mardocx_convert import convert
from mardocx_markdown import heading_warnings, prepare_images, quote_warnings
from mardocx_project import resolve_project_root
from mardocx_reference import create_reference, find_pandoc, update_reference
from mardocx import main as cli_main


class MardocxTests(unittest.TestCase):
    def config(self, root: Path):
        return load_config(SKILL_DIR, root)

    def test_heading_warnings_and_project_resolution(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            nested = root / "docs" / "drafts"
            nested.mkdir(parents=True)
            (root / "mardocx.yaml").write_text("{}", encoding="utf-8")
            markdown = nested / "report.md"
            markdown.write_text("## Starts late\n\n#### Too deep\n", encoding="utf-8")
            self.assertEqual(resolve_project_root(markdown, None), root)
            warnings, levels = heading_warnings(markdown)
            self.assertEqual(levels, [2, 4])
            self.assertEqual(len(warnings), 3)

    def test_reference_creation_and_explicit_update(self):
        try:
            find_pandoc()
        except Exception as exc:
            self.skipTest(str(exc))
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            config = self.config(root)
            reference = root / "reference.docx"
            create_reference(reference, config)
            before = reference.read_bytes()
            document = Document(reference)
            self.assertEqual(round(document.styles["Normal"].font.size.pt), 12)
            self.assertEqual(round(document.styles["Heading 1"].font.size.pt), 24)
            self.assertEqual(round(document.styles["Image Caption"].font.size.pt), 9)
            (root / "mardocx.yaml").write_text("typography:\n  body:\n    size_pt: 11\n", encoding="utf-8")
            update_reference(reference, self.config(root))
            self.assertNotEqual(before, reference.read_bytes())
            self.assertEqual(round(Document(reference).styles["Normal"].font.size.pt), 11)

    def test_local_image_rewrite(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            Image.new("RGB", (800, 400)).save(root / "wide.png")
            Image.new("RGB", (400, 800)).save(root / "tall.png")
            source = root / "input.md"
            source.write_text("![Wide](wide.png)\n![Tall](tall.png)\n![Remote](https://example.com/a.png)\n", encoding="utf-8")
            target = root / "prepared.md"
            warnings = prepare_images(source, target, self.config(root))
            prepared = target.read_text(encoding="utf-8")
            self.assertIn("{width=8cm}", prepared)
            self.assertIn("{width=6cm}", prepared)
            self.assertEqual(len(warnings), 1)

    def test_conversion_reuses_existing_reference(self):
        try:
            find_pandoc()
        except Exception as exc:
            self.skipTest(str(exc))
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            Image.new("RGB", (800, 400)).save(root / "wide.png")
            source = root / "report.md"
            source.write_text("# Chapter\n\n## Section\n\n![Figure one](wide.png)\n", encoding="utf-8")
            config = self.config(root)
            reference = root / "reference.docx"
            create_reference(reference, config)
            before = reference.read_bytes()
            output = root / "report.docx"
            messages = convert(source, root, output, config)
            self.assertEqual(messages, [])
            self.assertTrue(output.exists())
            self.assertEqual(before, reference.read_bytes())
            styles = {paragraph.style.name for paragraph in Document(output).paragraphs}
            self.assertTrue({"Heading 1", "Heading 2", "Image Caption"}.issubset(styles))

    def test_chinese_soft_breaks_and_ascii_quotes(self):
        try:
            find_pandoc()
        except Exception as exc:
            self.skipTest(str(exc))
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "repro.md"
            source.write_text(
                "中文第一行\n中文第二行\n\n他说：\"你好\"。\n",
                encoding="utf-8",
            )
            output = root / "repro.docx"
            convert(source, root, output, self.config(root))
            paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertEqual(paragraphs[0], "中文第一行中文第二行")
            self.assertEqual(paragraphs[1], "他说：“你好”。")
            self.assertEqual(
                [ord(character) for character in paragraphs[1] if character in "“”"],
                [0x201C, 0x201D],
            )

    def test_double_right_quotes_warn_without_rewriting_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "quotes.md"
            original = "他说：”你好”。\n```\n代码中的 ”忽略”\n```\n"
            source.write_text(original, encoding="utf-8")
            warnings = quote_warnings(source)
            self.assertEqual(len(warnings), 1)
            self.assertIn("possible double-right quotation mark", warnings[0])
            captured = StringIO()
            with redirect_stdout(captured):
                self.assertEqual(cli_main(["check", str(source)]), 0)
            self.assertIn("possible double-right quotation mark", captured.getvalue())
            conversion_warnings = convert(source, root, root / "quotes.docx", self.config(root))
            self.assertTrue(any("possible double-right quotation mark" in warning for warning in conversion_warnings))
            self.assertEqual(source.read_text(encoding="utf-8"), original)


if __name__ == "__main__":
    unittest.main()
